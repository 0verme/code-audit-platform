from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def detect_header_row(ws, header_aliases: dict[str, set[str]], normalize_value) -> tuple[int, dict[str, int]]:
    for row_index, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True), start=1):
        mapping = {}
        for col_index, value in enumerate(row, start=1):
            cell_text = normalize_value(value)
            for field_name, aliases in header_aliases.items():
                if cell_text in aliases:
                    mapping[field_name] = col_index
        if 'target_table' in mapping and 'target_column' in mapping and 'source_table' in mapping and 'source_column' in mapping:
            return row_index, mapping
    raise ValueError('鏈湪 Excel 涓瘑鍒埌鈥滅洰鏍囪〃/鐩爣瀛楁/婧愯〃/婧愬瓧娈碘€濊〃澶淬€?')


def load_lineage_edges_from_xlsx(
    xlsx_path: str | Path,
    *,
    detect_header_row_func,
    normalize_value,
) -> list[dict]:
    workbook = load_workbook(str(xlsx_path), read_only=True, data_only=True)
    try:
        last_error = None
        for worksheet in workbook.worksheets:
            try:
                header_row, columns = detect_header_row_func(worksheet)
            except ValueError as exc:
                last_error = exc
                continue
            rows = []
            for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
                source_table = normalize_value(values[columns['source_table'] - 1] if columns.get('source_table') else '')
                source_column = normalize_value(values[columns['source_column'] - 1] if columns.get('source_column') else '')
                target_table = normalize_value(values[columns['target_table'] - 1] if columns.get('target_table') else '')
                target_column = normalize_value(values[columns['target_column'] - 1] if columns.get('target_column') else '')
                if not (source_table and source_column and target_table and target_column):
                    continue
                rows.append({
                    'source_system': normalize_value(values[columns['source_system'] - 1] if columns.get('source_system') else ''),
                    'source_schema': normalize_value(values[columns['source_schema'] - 1] if columns.get('source_schema') else ''),
                    'source_table': source_table,
                    'source_column': source_column,
                    'target_system': normalize_value(values[columns['target_system'] - 1] if columns.get('target_system') else ''),
                    'target_schema': normalize_value(values[columns['target_schema'] - 1] if columns.get('target_schema') else ''),
                    'target_table': target_table,
                    'target_column': target_column,
                })
            return rows
        raise last_error or ValueError('鏈湪 Excel 涓瘑鍒埌鏈夋晥琛€缂樻暟鎹€?')
    finally:
        workbook.close()
