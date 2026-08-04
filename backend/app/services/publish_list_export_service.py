from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.publish_list_service import get_publish_list


EXPORT_HEADERS = ("需求 ID", "需求类型", "开发人员", "需求标题", "状态", "计划发布时间")
EXPORTABLE_STATUSES = frozenset(
    {"pending", "reviewing", "passed", "rejected", "launched", "processing", "completed", "unknown"}
)
_COLUMN_WIDTHS = (18, 14, 16, 42, 16, 22)
_STATUS_LABELS = {
    "pending": "待审核",
    "reviewing": "审核中",
    "passed": "等待上线",
    "rejected": "已驳回",
    "launched": "已上线",
    "processing": "处理中",
    "completed": "处理完成",
    "unknown": "未知状态",
}


def _planned_at(item: dict[str, Any], fallback_date: date) -> datetime:
    raw_date = str(item.get("date") or fallback_date.isoformat()).strip()[:10]
    try:
        planned_date = date.fromisoformat(raw_date)
    except ValueError:
        planned_date = fallback_date

    raw_time = str(item.get("time") or "").strip()
    if not raw_time:
        return datetime.combine(planned_date, time.min)
    try:
        planned_time = time.fromisoformat(raw_time)
    except ValueError:
        return datetime.combine(planned_date, time.min)
    return datetime.combine(planned_date, planned_time.replace(tzinfo=None))


def _write_text(cell, value: Any) -> None:
    cell.value = str(value or "")
    cell.data_type = "s"


def _status_label(item: dict[str, Any]) -> str:
    label = str(item.get("statusLabel") or "").strip()
    if any("\u4e00" <= character <= "\u9fff" for character in label):
        return label
    return _STATUS_LABELS.get(str(item.get("status") or "unknown"), "未知状态")


def _filter_items(
    items: Iterable[dict[str, Any]],
    *,
    status: str | None,
    requirement_types: Iterable[str] | None,
) -> list[dict[str, Any]]:
    selected_types = {value for value in (requirement_types or ()) if value}
    return [
        item
        for item in items
        if (status is None or item.get("status") == status)
        and (not selected_types or item.get("type") in selected_types)
    ]


def build_publish_list_workbook(
    selected_date: date,
    *,
    status: str | None = None,
    requirement_types: Iterable[str] | None = None,
    profile: Any = None,
    runner: Any = None,
) -> BytesIO:
    if status is not None and status not in EXPORTABLE_STATUSES:
        raise ValueError("unsupported publish-list status")

    payload = get_publish_list(selected_date, profile=profile, runner=runner, limit=None)
    items = _filter_items(payload["items"], status=status, requirement_types=requirement_types)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "上线清单"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    sheet.append(EXPORT_HEADERS)
    header_fill = PatternFill(fill_type="solid", fgColor="2563EB")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 24

    for item in items:
        row_index = sheet.max_row + 1
        values = (
            item.get("id"),
            item.get("type"),
            item.get("owner"),
            item.get("title"),
            _status_label(item),
        )
        for column_index, value in enumerate(values, start=1):
            _write_text(sheet.cell(row=row_index, column=column_index), value)
        planned_at_cell = sheet.cell(row=row_index, column=6, value=_planned_at(item, selected_date))
        planned_at_cell.number_format = "yyyy-mm-dd hh:mm"
        for cell in sheet[row_index]:
            cell.alignment = Alignment(vertical="center", wrap_text=cell.column == 4)

    sheet.auto_filter.ref = f"A1:F{sheet.max_row}"
    for index, width in enumerate(_COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output
