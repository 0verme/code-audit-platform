import sqlite3
import sys
import tempfile
import unittest
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook


BACKEND_DIR = Path(__file__).resolve().parents[1] / "backend"
SVN_CHECK_DIR = BACKEND_DIR / "svn_check"
for path in (BACKEND_DIR, SVN_CHECK_DIR):
    if str(path) not in sys.path:
        sys.path.insert(0, str(path))

from lineage import mapping_compat as mapping_sqlite  # noqa: E402
from lineage import mapping_compat as new_mapping  # noqa: E402


def alias(field_name: str) -> str:
    return next(iter(mapping_sqlite.HEADER_ALIASES[field_name]))


def build_workbook(path: Path, sheets: list[dict]) -> None:
    workbook = Workbook()
    first = True
    for sheet_spec in sheets:
        if first:
            worksheet = workbook.active
            worksheet.title = sheet_spec["title"]
            first = False
        else:
            worksheet = workbook.create_sheet(sheet_spec["title"])
        for row in sheet_spec["rows"]:
            worksheet.append(row)
    workbook.save(path)
    workbook.close()


class MappingSqliteContractTests(unittest.TestCase):
    @contextmanager
    def tracked_sqlite_connections(self):
        original_connect = sqlite3.connect
        opened = []

        def connect(*args, **kwargs):
            conn = original_connect(*args, **kwargs)
            opened.append(conn)
            return conn

        with patch.object(mapping_sqlite.sqlite3, "connect", side_effect=connect):
            yield
        for conn in reversed(opened):
            try:
                conn.close()
            except Exception:
                pass

    def test_registered_result_table_helpers_keep_current_normalization_behavior(self):
        rows = [
            (" dm.table_a extra ",),
            ("DM.TABLE_A",),
            ("  ",),
            (None,),
            ("ods.table_b",),
            (" ODS.TABLE_B trailing ",),
        ]

        with patch.object(mapping_sqlite, "select_sql_with_profile", return_value=rows):
            result_tables = mapping_sqlite.load_registered_result_tables(profile="czcb")

        self.assertEqual(result_tables, {"DM.TABLE_A", "ODS.TABLE_B"})
        self.assertEqual(mapping_sqlite.normalize_registered_table_name(" dm.table_c extra "), "DM.TABLE_C")
        self.assertEqual(
            mapping_sqlite.filter_registered_result_nodes(
                [
                    ("dm", "table_a", "col_a"),
                    ("ODS", "TABLE_B", "col_b"),
                    ("dm", "table_x", "col_x"),
                ],
                result_tables=result_tables,
            ),
            [("dm", "table_a", "col_a"), ("ODS", "TABLE_B", "col_b")],
        )

    def test_load_mapping_meta_and_status_return_empty_state_when_files_are_missing(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            db_path = tmp_path / "missing.db"
            xlsx_path = tmp_path / "missing.xlsx"

            self.assertEqual(mapping_sqlite.load_mapping_meta(db_path), {})
            self.assertEqual(
                mapping_sqlite.get_mapping_db_status(db_path, xlsx_path),
                {
                    "db_exists": False,
                    "xlsx_exists": False,
                    "is_fresh": False,
                    "db_path": str(db_path),
                    "xlsx_path": str(xlsx_path),
                    "meta": {},
                },
            )

    def test_recreate_mapping_sqlite_builds_cache_tables_and_meta_contract(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            xlsx_path = tmp_path / "mapping.xlsx"
            db_path = tmp_path / "mapping.db"
            build_workbook(
                xlsx_path,
                [
                    {
                        "title": "lineage",
                        "rows": [
                            [alias("target_table"), alias("target_column"), alias("source_table"), alias("source_column")],
                            ["DM.TGT_A", "COL_A", "ODS.SRC_A", "SRC_COL_A"],
                            ["DM.TGT_B", "COL_B", "ODS.SRC_B", "SRC_COL_B"],
                        ],
                    }
                ],
            )

            with self.tracked_sqlite_connections():
                result = mapping_sqlite.recreate_mapping_sqlite(xlsx_path, db_path)
                meta = mapping_sqlite.load_mapping_meta(db_path)
                status = mapping_sqlite.get_mapping_db_status(db_path, xlsx_path)

                self.assertEqual(result["db_path"], str(db_path))
                self.assertEqual(result["xlsx_path"], str(xlsx_path))
                self.assertEqual(result["edge_count"], 2)
                self.assertEqual(meta["source_xlsx_path"], str(xlsx_path))
                self.assertEqual(meta["source_xlsx_mtime_ns"], str(xlsx_path.stat().st_mtime_ns))
                self.assertEqual(meta["source_xlsx_size"], str(xlsx_path.stat().st_size))
                self.assertEqual(meta["edge_count"], "2")
                self.assertTrue(meta["imported_at"])
                self.assertTrue(status["db_exists"])
                self.assertTrue(status["xlsx_exists"])
                self.assertTrue(status["is_fresh"])

                with sqlite3.connect(db_path) as conn:
                    tables = {row[0] for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                    edge_rows = conn.execute(
                        "SELECT source_schema, source_table, source_column, target_schema, target_table, target_column "
                        "FROM lineage_edge ORDER BY id"
                    ).fetchall()

            self.assertIn("lineage_edge", tables)
            self.assertIn("lineage_meta", tables)
            self.assertEqual(
                edge_rows,
                [
                    ("ODS", "SRC_A", "SRC_COL_A", "DM", "TGT_A", "COL_A"),
                    ("ODS", "SRC_B", "SRC_COL_B", "DM", "TGT_B", "COL_B"),
                ],
            )

    def test_get_mapping_db_status_marks_cache_stale_when_meta_mismatch_exists(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            xlsx_path = tmp_path / "mapping.xlsx"
            db_path = tmp_path / "mapping.db"
            build_workbook(
                xlsx_path,
                [{"title": "lineage", "rows": [[alias("target_table"), alias("target_column"), alias("source_table"), alias("source_column")], ["DM.TGT_A", "COL_A", "ODS.SRC_A", "SRC_COL_A"]]}],
            )
            with self.tracked_sqlite_connections():
                mapping_sqlite.recreate_mapping_sqlite(xlsx_path, db_path)

                with sqlite3.connect(db_path) as conn:
                    conn.execute("UPDATE lineage_meta SET value = '0' WHERE key = 'source_xlsx_size'")
                    conn.commit()

                status = mapping_sqlite.get_mapping_db_status(db_path, xlsx_path)
                self.assertTrue(status["db_exists"])
                self.assertTrue(status["xlsx_exists"])
                self.assertFalse(status["is_fresh"])
                self.assertEqual(status["meta"]["source_xlsx_size"], "0")

    def test_detect_header_row_and_load_lineage_edges_from_xlsx_keep_alias_and_sheet_selection_behavior(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            xlsx_path = Path(tmpdir) / "mapping.xlsx"
            build_workbook(
                xlsx_path,
                [
                    {"title": "ignored", "rows": [["not", "a", "header"]]},
                    {
                        "title": "valid",
                        "rows": [
                            ["note", "", "", ""],
                            [alias("target_system"), alias("target_schema"), alias("target_table"), alias("target_column"),
                             alias("source_system"), alias("source_schema"), alias("source_table"), alias("source_column")],
                            ["dm", "dm", "target_a", "col_a", "ods", "ods", "source_a", "src_col_a"],
                            ["", "", "", "", "", "", "", ""],
                            ["dm", "", "target_b", "col_b", "ods", "", "source_b", "src_col_b"],
                        ],
                    },
                ],
            )

            from openpyxl import load_workbook

            loaded = load_workbook(xlsx_path, read_only=True, data_only=True)
            try:
                header_row, columns = mapping_sqlite.detect_header_row(loaded["valid"])
            finally:
                loaded.close()

            edges = mapping_sqlite.load_lineage_edges_from_xlsx(xlsx_path)

            self.assertEqual(header_row, 2)
            self.assertIn("target_table", columns)
            self.assertEqual(
                edges,
                [
                    {
                        "source_system": "ods",
                        "source_schema": "ods",
                        "source_table": "source_a",
                        "source_column": "src_col_a",
                        "target_system": "dm",
                        "target_schema": "dm",
                        "target_table": "target_a",
                        "target_column": "col_a",
                    },
                    {
                        "source_system": "ods",
                        "source_schema": "",
                        "source_table": "source_b",
                        "source_column": "src_col_b",
                        "target_system": "dm",
                        "target_schema": "",
                        "target_table": "target_b",
                        "target_column": "col_b",
                    },
                ],
            )

    def test_identifier_and_traversal_helpers_keep_current_behavior(self):
        self.assertEqual(mapping_sqlite.parse_input_table_name(" dm.table_a "), ("DM", "TABLE_A"))
        self.assertEqual(mapping_sqlite.parse_input_table_name("table_a"), ("", "TABLE_A"))
        self.assertEqual(mapping_sqlite.normalize_identifier("", "dm.table_a", " col_a "), ("DM", "TABLE_A", "COL_A"))
        self.assertEqual(mapping_sqlite.compact_identifier(("DM", "TABLE_A", "COL_A")), "DM.TABLE_A.COL_A")

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            db_path = tmp_path / "mapping.db"
            with self.tracked_sqlite_connections():
                with sqlite3.connect(db_path) as conn:
                    conn.execute(
                        """
                        CREATE TABLE lineage_edge (
                            id INTEGER PRIMARY KEY AUTOINCREMENT,
                            source_system TEXT NOT NULL,
                            source_schema TEXT NOT NULL,
                            source_table TEXT NOT NULL,
                            source_column TEXT NOT NULL,
                            target_system TEXT NOT NULL,
                            target_schema TEXT NOT NULL,
                            target_table TEXT NOT NULL,
                            target_column TEXT NOT NULL
                        )
                        """
                    )
                    conn.executemany(
                        """
                        INSERT INTO lineage_edge (
                            source_system, source_schema, source_table, source_column,
                            target_system, target_schema, target_table, target_column
                        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                        """,
                        [
                            ("ODS", "DM", "TABLE_A", "COL_A", "DM", "DM", "TABLE_B", "COL_B"),
                            ("ODS", "DM", "TABLE_A", "COL_A", "DM", "DM", "TABLE_C", "COL_C"),
                            ("ODS", "DM", "TABLE_B", "COL_B", "DM", "DM", "TABLE_D", "COL_D"),
                            ("ODS", "DM", "TABLE_C", "COL_C", "DM", "DM", "TABLE_D", "COL_D"),
                        ],
                    )
                    conn.commit()

                start_nodes = mapping_sqlite.find_start_nodes_in_sqlite("dm.table_a", "col_a", db_path)
                relation_rows, ordered_nodes = mapping_sqlite.walk_downstream_in_sqlite(start_nodes, db_path, max_depth=None)

                self.assertEqual(start_nodes, [("DM", "TABLE_A", "COL_A")])
                self.assertEqual(
                    relation_rows,
                    [
                        [1, "DM.TABLE_A.COL_A", "DM.TABLE_B.COL_B"],
                        [1, "DM.TABLE_A.COL_A", "DM.TABLE_C.COL_C"],
                        [2, "DM.TABLE_B.COL_B", "DM.TABLE_D.COL_D"],
                        [2, "DM.TABLE_C.COL_C", "DM.TABLE_D.COL_D"],
                    ],
                )
                self.assertEqual(
                    ordered_nodes,
                    [("DM", "TABLE_B", "COL_B"), ("DM", "TABLE_C", "COL_C"), ("DM", "TABLE_D", "COL_D")],
                )

    def test_import_paths_remain_compatible(self):
        from shared.lineage.mapping_sqlite import load_registered_result_tables  # noqa: PLC0415
        from shared.lineage import mapping_sqlite as legacy_mapping  # noqa: PLC0415
        from shared.lineage import mapping_sqlite as imported_module  # noqa: PLC0415

        self.assertIs(load_registered_result_tables, legacy_mapping.load_registered_result_tables)
        self.assertIs(imported_module, legacy_mapping)
        self.assertIs(mapping_sqlite.normalize_identifier, new_mapping.normalize_identifier)
        self.assertIs(mapping_sqlite.recreate_mapping_sqlite, new_mapping.recreate_mapping_sqlite)
        self.assertEqual(mapping_sqlite.MAPPING_DB_PATH, new_mapping.MAPPING_DB_PATH)
        self.assertEqual(mapping_sqlite.MAPPING_XLSX_PATH, new_mapping.MAPPING_XLSX_PATH)


if __name__ == "__main__":
    unittest.main()
