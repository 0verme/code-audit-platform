import sys
import unittest
from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


BACKEND_DIR = Path(__file__).resolve().parents[1] / "backend"
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.services.publish_list_service import PublishListConfigurationError, get_publish_list  # noqa: E402
from app.services.publish_list_export_service import build_publish_list_workbook  # noqa: E402


class FakeProfile:
    config = {
        "schema": "dwp",
        "metadata": {
            "tables": {"publish_list": "p_publishlist"},
            "columns": {
                "publish_list": {
                    "id": "req_id",
                    "title": "req_title",
                    "status": "publish_status",
                    "date": "publish_date",
                    "owner": "developer",
                }
            },
        },
    }


class AutoMappedProfile:
    config = {"schema": "dwp"}


class FakeRunner:
    def __init__(self, columns, rows):
        self.columns = columns
        self.rows = rows
        self.queries = []

    def query_all(self, sql, params=None):
        self.queries.append((sql, params))
        if "information_schema.columns" in sql:
            return [{"column_name": column} for column in self.columns]
        return self.rows


class PublishListServiceTests(unittest.TestCase):
    def test_auto_maps_task_columns_from_test_table_schema(self):
        columns = ["taskid", "tasktype", "taskuer", "tasktitle", "taskstatus", "taskdate", "tasksource", "taskremark"]
        runner = FakeRunner(columns, [])

        get_publish_list(date(2026, 7, 22), profile=AutoMappedProfile(), runner=runner)

        sql = runner.queries[1][0]
        self.assertIn("taskid AS id", sql)
        self.assertIn("taskuer AS owner", sql)
        self.assertIn("taskremark AS remark", sql)
        self.assertIn("WHERE taskdate = ?", sql)
        self.assertNotIn("CAST(", sql)
        self.assertIn("ORDER BY taskid LIMIT 1000", sql)

    def test_maps_profile_columns_and_normalizes_rows(self):
        runner = FakeRunner(
            ["req_id", "req_title", "publish_status", "publish_date", "developer", "taskremark"],
            [{
                "id": "REQ-1", "type": None, "owner": "张三", "title": "上线需求",
                "status": "已通过", "date": date(2026, 7, 22), "time": None,
                "source": None, "remark": "测试备注", "detail_url": "javascript:alert(1)",
            }],
        )
        payload = get_publish_list(date(2026, 7, 22), profile=FakeProfile(), runner=runner)

        self.assertEqual(payload["items"][0]["status"], "passed")
        self.assertEqual(payload["items"][0]["date"], "2026-07-22")
        self.assertEqual(payload["items"][0]["detailUrl"], "")
        self.assertEqual(payload["items"][0]["remark"], "测试备注")
        self.assertEqual(payload["summary"]["passed"], 1)
        self.assertIn("FROM dwp.p_publishlist", runner.queries[1][0])
        self.assertEqual(runner.queries[1][1], ("2026-07-22",))

    def test_rejects_incomplete_field_mapping(self):
        runner = FakeRunner(["req_id", "req_title"], [])
        with self.assertRaises(PublishListConfigurationError):
            get_publish_list(date(2026, 7, 22), profile=FakeProfile(), runner=runner)

    def test_uses_http_detail_field_as_new_window_link(self):
        runner = FakeRunner(
            ["req_id", "req_title", "publish_status", "publish_date", "developer", "taskremark"],
            [{
                "id": "REQ-1", "type": None, "owner": None, "title": "演示",
                "status": "审核中", "date": date(2026, 7, 22), "time": None,
                "source": None, "remark": "https://example.com/detail/REQ-1", "detail_url": None,
            }],
        )

        payload = get_publish_list(date(2026, 7, 22), profile=FakeProfile(), runner=runner)

        self.assertEqual(payload["items"][0]["detailUrl"], "https://example.com/detail/REQ-1")

    def test_normalizes_business_status_codes(self):
        rows = []
        for index, status in enumerate(("生产待审", "运行待审", "开发待审", "待审核", "待审查", "等待上线"), start=1):
            rows.append({
                "id": f"REQ-{index}", "type": None, "owner": None, "title": "演示",
                "status": status, "date": date(2026, 7, 22), "time": None,
                "source": None, "remark": None, "detail_url": None,
            })
        runner = FakeRunner(
            ["req_id", "req_title", "publish_status", "publish_date", "developer"],
            rows,
        )

        payload = get_publish_list(date(2026, 7, 22), profile=FakeProfile(), runner=runner)

        statuses_by_id = {item["id"]: item["status"] for item in payload["items"]}
        self.assertEqual(statuses_by_id, {
            "REQ-1": "reviewing", "REQ-2": "reviewing", "REQ-3": "reviewing",
            "REQ-4": "reviewing", "REQ-5": "reviewing", "REQ-6": "passed",
        })
        self.assertEqual(
            [item["statusLabel"] for item in payload["items"] if item["status"] == "reviewing"],
            ["审核中"] * 5,
        )
        self.assertEqual(payload["summary"]["pending"], 0)
        self.assertEqual(payload["summary"]["reviewing"], 5)
        self.assertEqual(payload["summary"]["passed"], 1)

    def test_sorts_by_business_status_then_requirement_type(self):
        rows = [
            {"id": "REQ-1", "type": "运行维护", "owner": None, "title": "演示", "status": "审核中", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-2", "type": "数据修改", "owner": None, "title": "演示", "status": "等待上线", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-3", "type": "开发维护", "owner": None, "title": "演示", "status": "等待上线", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-4", "type": "数据修改", "owner": None, "title": "演示", "status": "开发待审", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-5", "type": "开发维护", "owner": None, "title": "演示", "status": "开发待审", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-6", "type": "开发维护", "owner": None, "title": "演示", "status": "已上线", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-7", "type": "数据修改", "owner": None, "title": "演示", "status": "处理完成", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-8", "type": "运行维护", "owner": None, "title": "演示", "status": "处理中", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
        ]
        runner = FakeRunner(
            ["req_id", "req_title", "publish_status", "publish_date", "developer", "requirement_type"],
            rows,
        )

        payload = get_publish_list(date(2026, 7, 22), profile=FakeProfile(), runner=runner)

        self.assertEqual([item["id"] for item in payload["items"]], ["REQ-6", "REQ-3", "REQ-2", "REQ-7", "REQ-5", "REQ-4", "REQ-1", "REQ-8"])

    def test_normalizes_processing_statuses_and_counts_them_separately(self):
        rows = [
            {"id": "REQ-1", "type": "开发维护", "owner": None, "title": "演示", "status": "处理中", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
            {"id": "REQ-2", "type": "运行维护", "owner": None, "title": "演示", "status": "处理完成", "date": date(2026, 7, 22), "time": None, "source": None, "remark": None, "detail_url": None},
        ]
        runner = FakeRunner(
            ["req_id", "req_title", "publish_status", "publish_date", "developer", "requirement_type"],
            rows,
        )

        payload = get_publish_list(date(2026, 7, 22), profile=FakeProfile(), runner=runner)

        self.assertEqual([item["status"] for item in payload["items"]], ["completed", "processing"])
        self.assertEqual(payload["summary"]["processing"], 1)
        self.assertEqual(payload["summary"]["completed"], 1)
        self.assertEqual(payload["summary"]["launched"], 0)

    def test_unlimited_query_omits_page_limit(self):
        runner = FakeRunner(
            ["req_id", "req_title", "publish_status", "publish_date", "developer"],
            [],
        )

        get_publish_list(date(2026, 8, 4), profile=FakeProfile(), runner=runner, limit=None)

        self.assertNotIn("LIMIT", runner.queries[1][0])


class PublishListExportTests(unittest.TestCase):
    def _runner(self, rows):
        return FakeRunner(
            [
                "req_id", "req_title", "publish_status", "publish_date", "developer",
                "requirement_type", "planned_publish_time",
            ],
            rows,
        )

    def test_builds_filtered_safe_workbook_with_typed_publish_time(self):
        rows = [
            {
                "id": "=1+1", "type": "开发维护", "owner": "+cmd", "title": "@上线标题",
                "status": "2", "date": date(2026, 8, 4), "time": time(9, 30),
                "source": None, "remark": None, "detail_url": None,
            },
            {
                "id": "REQ-2", "type": "运行维护", "owner": "李四", "title": "其他标题",
                "status": "已上线", "date": date(2026, 8, 4), "time": time(10, 0),
                "source": None, "remark": None, "detail_url": None,
            },
        ]

        output = build_publish_list_workbook(
            date(2026, 8, 4),
            status="passed",
            requirement_types=["开发维护"],
            profile=FakeProfile(),
            runner=self._runner(rows),
        )
        sheet = load_workbook(output)["上线清单"]

        self.assertEqual([cell.value for cell in sheet[1]], ["需求 ID", "需求类型", "开发人员", "需求标题", "状态", "计划发布时间"])
        self.assertEqual(sheet.max_row, 2)
        self.assertEqual([sheet.cell(2, column).data_type for column in range(1, 6)], ["s"] * 5)
        self.assertEqual(sheet["A2"].value, "=1+1")
        self.assertEqual(sheet["C2"].value, "+cmd")
        self.assertEqual(sheet["D2"].value, "@上线标题")
        self.assertEqual(sheet["E2"].value, "等待上线")
        self.assertEqual(sheet["F2"].value, datetime(2026, 8, 4, 9, 30))
        self.assertEqual(sheet["F2"].number_format, "yyyy-mm-dd hh:mm")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, "A1:F2")

    def test_empty_filter_creates_header_only_workbook(self):
        output = build_publish_list_workbook(
            date(2026, 8, 4),
            requirement_types=["数据修改"],
            profile=FakeProfile(),
            runner=self._runner([]),
        )

        sheet = load_workbook(output)["上线清单"]
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet.auto_filter.ref, "A1:F1")

    def test_rejects_unknown_status_before_querying(self):
        with self.assertRaises(ValueError):
            build_publish_list_workbook(date(2026, 8, 4), status="bogus")

    def test_export_route_returns_xlsx_attachment(self):
        import app as app_module  # noqa: E402

        workbook = Workbook()
        payload = BytesIO()
        workbook.save(payload)
        payload.seek(0)
        with patch("app.routes.publish_list.build_publish_list_workbook", return_value=payload) as builder:
            response = app_module.create_app().test_client().get(
                "/api/publish-list/export?date=2026-08-04&status=passed&type=%E5%BC%80%E5%8F%91%E7%BB%B4%E6%8A%A4"
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment", response.headers["Content-Disposition"])
        self.assertIn("filename*=UTF-8''", response.headers["Content-Disposition"])
        builder.assert_called_once_with(
            date(2026, 8, 4), status="passed", requirement_types=["开发维护"]
        )
        response.close()

    def test_export_route_rejects_invalid_date_and_status(self):
        import app as app_module  # noqa: E402

        client = app_module.create_app().test_client()
        invalid_date = client.get("/api/publish-list/export?date=08-04-2026")
        invalid_status = client.get("/api/publish-list/export?date=2026-08-04&status=bogus")

        self.assertEqual(invalid_date.status_code, 400)
        self.assertEqual(invalid_status.status_code, 400)
        invalid_date.close()
        invalid_status.close()


if __name__ == "__main__":
    unittest.main()
